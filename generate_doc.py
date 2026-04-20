"""Generate .docx report for AI Service assignment submission."""
import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# ── Page margins
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(13)

def add_heading_custom(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h

def add_code_block(code_text, font_size=8):
    """Add code with monospace font."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(font_size)
    return p

def add_image_or_placeholder(image_path, caption, width=Inches(5.5)):
    """Add image if exists, else add placeholder comment."""
    if os.path.exists(image_path):
        doc.add_picture(image_path, width=width)
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p = doc.add_paragraph()
        run = p.add_run(f'[CHÈN ẢNH: {caption}]')
        run.font.color.rgb = RGBColor(255, 0, 0)
        run.bold = True
        run.font.size = Pt(12)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Caption
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(10)

# ════════════════════════════════════════════════
# 1. TRANG BÌA
# ════════════════════════════════════════════════
for _ in range(4):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('TRƯỜNG ĐẠI HỌC ...')
run.font.size = Pt(16)
run.bold = True
run.font.name = 'Times New Roman'

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('KHOA CÔNG NGHỆ THÔNG TIN')
run.font.size = Pt(14)
run.bold = True

for _ in range(3):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('BÁO CÁO BÀI TẬP')
run.font.size = Pt(20)
run.bold = True
run.font.color.rgb = RGBColor(0, 51, 153)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('AI SERVICE')
run.font.size = Pt(24)
run.bold = True
run.font.color.rgb = RGBColor(0, 51, 153)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Xây dựng mô hình RNN/LSTM/BiLSTM + Knowledge Graph + RAG\ncho hệ thống E-commerce')
run.font.size = Pt(14)
run.italic = True

for _ in range(4):
    doc.add_paragraph()

info_lines = [
    'Lớp: [TÊN LỚP]',
    'Nhóm: [SỐ NHÓM]',
    'Thành viên: [HỌ TÊN]',
    'GVHD: [TÊN GIẢNG VIÊN]',
    'Ngày nộp: 20/04/2026',
]
for line in info_lines:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(line)
    run.font.size = Pt(13)

doc.add_page_break()

# ════════════════════════════════════════════════
# 2. MÔ TẢ AI SERVICE
# ════════════════════════════════════════════════
add_heading_custom('1. MÔ TẢ AI SERVICE', level=1)

doc.add_paragraph(
    'AI New Service là một microservice AI được xây dựng và tích hợp vào hệ thống '
    'E-commerce microservices (Django + FastAPI). Service này cung cấp các chức năng:'
)

features = [
    'Bảng UserBehaviorData: Lưu trữ hành vi người dùng với 8 loại action (view, click, add_to_cart, purchase, search, wishlist, remove_from_cart, review)',
    '3 mô hình Deep Learning: RNN, LSTM, BiLSTM để dự đoán hành vi tiếp theo và phân loại hành vi người dùng',
    'Knowledge Base Graph (Neo4j): Xây dựng đồ thị tri thức biểu diễn quan hệ User-Product-Category',
    'RAG Chat Engine: Chatbot thông minh kết hợp Knowledge Graph + ChromaDB Vector Store',
    'Tích hợp Frontend: Tab mới trên giao diện e-commerce với danh sách recommend và chat AI',
]
for f in features:
    doc.add_paragraph(f, style='List Bullet')

add_heading_custom('Kiến trúc hệ thống', level=2)
doc.add_paragraph(
    'Service được triển khai dưới dạng Docker container, giao tiếp qua REST API với các service khác '
    '(auth-service, product-service, cart-service) thông qua API Gateway. '
    'Sử dụng PostgreSQL riêng cho behavior data, Neo4j cho Knowledge Graph, '
    'ChromaDB cho vector embeddings, và Redis cho caching.'
)

add_heading_custom('Công nghệ sử dụng', level=2)
techs = [
    ('Backend', 'FastAPI (Python 3.11)'),
    ('Deep Learning', 'PyTorch (RNN, LSTM, BiLSTM)'),
    ('Graph Database', 'Neo4j 5 Community Edition'),
    ('Vector Store', 'ChromaDB + sentence-transformers'),
    ('Database', 'PostgreSQL 15'),
    ('Frontend', 'React (Vite)'),
    ('Containerization', 'Docker Compose'),
]
table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
hdr[0].text = 'Thành phần'
hdr[1].text = 'Công nghệ'
for name, tech in techs:
    row = table.add_row().cells
    row[0].text = name
    row[1].text = tech

doc.add_page_break()

# ════════════════════════════════════════════════
# 3. CÂU 1 — DATA 500 USER + 8 BEHAVIORS (Copy 20 dòng)
# ════════════════════════════════════════════════
add_heading_custom('2. CÂU 1: Sinh tập dữ liệu data_user500 (500 user + 8 behaviors)', level=1)

add_heading_custom('2.1. Mô tả', level=2)
doc.add_paragraph(
    'Tập dữ liệu được sinh tự động khi ai-new-service khởi động. '
    'Hệ thống thực hiện 3 bước:\n'
    '1. Seed 500 user THẬT vào auth-service database thông qua API /api/auth/register/\n'
    '2. Fetch tất cả product ID thật từ product-service database\n'
    '3. Sinh behavior data với các pattern thực tế (funnel: view→click→add_to_cart→purchase)'
)

add_heading_custom('2.2. Bảng UserBehaviorData', level=2)
doc.add_paragraph('Cấu trúc bảng trong PostgreSQL:')

table = doc.add_table(rows=1, cols=4)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Column'
hdr[1].text = 'Type'
hdr[2].text = 'Nullable'
hdr[3].text = 'Description'
cols_data = [
    ('id', 'INTEGER (PK)', 'NOT NULL', 'Auto increment'),
    ('user_id', 'VARCHAR(255)', 'NOT NULL', 'UUID từ auth-service'),
    ('product_id', 'VARCHAR(255)', 'NOT NULL', 'UUID từ product-service'),
    ('action', 'VARCHAR(50)', 'NOT NULL', '8 loại: view, click, add_to_cart, purchase, search, wishlist, remove_from_cart, review'),
    ('timestamp', 'DATETIME', 'NOT NULL', 'Thời điểm hành vi'),
]
for c in cols_data:
    row = table.add_row().cells
    for i, val in enumerate(c):
        row[i].text = val

add_heading_custom('2.3. Code sinh dữ liệu', level=2)
doc.add_paragraph('File: services/ai-new-service/app/data_generator.py')

code_data_gen = '''ACTIONS = ["view", "click", "add_to_cart", "purchase", 
           "search", "wishlist", "remove_from_cart", "review"]

# 4 user behavior profiles
PROFILES = [
    # Browser: mostly views and clicks, rarely buys (40%)
    {"funnel_prob": 0.1, "actions_dist": [40, 30, 10, 2, 10, 5, 2, 1]},
    # Active Shopper: full funnel view→click→cart→purchase (25%)
    {"funnel_prob": 0.7, "actions_dist": [20, 20, 25, 15, 5, 5, 5, 5]},
    # Researcher: searches a lot, compares (20%)
    {"funnel_prob": 0.3, "actions_dist": [15, 15, 10, 5, 35, 10, 5, 5]},
    # Impulse Buyer: quick view→purchase (15%)
    {"funnel_prob": 0.5, "actions_dist": [25, 10, 15, 20, 5, 5, 5, 15]},
]'''
add_code_block(code_data_gen)

add_heading_custom('2.4. Mẫu 20 dòng dữ liệu', level=2)
doc.add_paragraph('Dưới đây là 20 dòng mẫu từ bảng user_behavior_data:')

table = doc.add_table(rows=1, cols=5)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'id'
hdr[1].text = 'user_id'
hdr[2].text = 'product_id'
hdr[3].text = 'action'
hdr[4].text = 'timestamp'

sample_actions = ['view', 'view', 'click', 'click', 'add_to_cart', 'purchase',
                  'view', 'search', 'click', 'view', 'wishlist', 'view',
                  'click', 'add_to_cart', 'view', 'review', 'remove_from_cart',
                  'view', 'click', 'search']
for i in range(20):
    row = table.add_row().cells
    row[0].text = str(i + 1)
    row[1].text = f'xxxxxxxx-xxxx-...-{i+1:04d}'
    row[2].text = f'yyyyyyyy-yyyy-...-{i+1:04d}'
    row[3].text = sample_actions[i]
    row[4].text = f'2026-0{(i%3)+1}-{(i%28)+1:02d} {(i*3)%24:02d}:{(i*7)%60:02d}:00'

p = doc.add_paragraph()
run = p.add_run('[GHI CHÚ: Thay bằng dữ liệu thật từ DB sau khi chạy docker compose up --build. '
                'Query: SELECT * FROM user_behavior_data LIMIT 20;]')
run.font.color.rgb = RGBColor(255, 0, 0)
run.italic = True
run.font.size = Pt(10)

doc.add_page_break()

# ════════════════════════════════════════════════
# 4. CÂU 2a — RNN, LSTM, BiLSTM + ĐÁNH GIÁ
# ════════════════════════════════════════════════
add_heading_custom('3. CÂU 2a: Xây dựng 3 mô hình RNN, LSTM, BiLSTM', level=1)

add_heading_custom('3.1. Kiến trúc 3 mô hình', level=2)

# RNN
add_heading_custom('3.1.1. RNN (Vanilla Recurrent Neural Network)', level=3)
doc.add_paragraph(
    'Mô hình RNN sử dụng 1-layer RNN cơ bản. Input là concatenation của '
    'Product Embedding (dim=64) và Action Embedding (dim=64). '
    'Output gồm 2 heads: phân loại action tiếp theo và dự đoán product tiếp theo.'
)
doc.add_paragraph('File: services/ai-new-service/app/models/behavior_models.py')

code_rnn = '''class BehaviorRNN(nn.Module):
    def __init__(self, num_users, num_products, num_actions, 
                 embedding_dim=64, hidden_dim=128):
        super().__init__()
        self.model_name = "RNN"
        self.product_emb = nn.Embedding(num_products, embedding_dim, padding_idx=0)
        self.action_emb = nn.Embedding(num_actions, embedding_dim, padding_idx=0)
        input_dim = embedding_dim * 2
        self.rnn = nn.RNN(input_dim, hidden_dim, num_layers=1, batch_first=True)
        self.dropout = nn.Dropout(0.2)
        self.fc_action = nn.Linear(hidden_dim, num_actions)
        self.fc_product = nn.Linear(hidden_dim, num_products)

    def forward(self, user_ids, product_ids, action_ids):
        p = self.product_emb(product_ids)
        a = self.action_emb(action_ids)
        x = torch.cat([p, a], dim=-1)
        out, _ = self.rnn(x)
        last = self.dropout(out[:, -1, :])
        return self.fc_action(last), self.fc_product(last)'''
add_code_block(code_rnn)

# LSTM
add_heading_custom('3.1.2. LSTM (Long Short-Term Memory)', level=3)
doc.add_paragraph(
    'Mô hình LSTM sử dụng 2-layer LSTM với dropout=0.2. '
    'LSTM có cơ chế cell state giúp học được long-term dependencies tốt hơn RNN.'
)

code_lstm = '''class BehaviorLSTM(nn.Module):
    def __init__(self, num_users, num_products, num_actions,
                 embedding_dim=64, hidden_dim=128):
        super().__init__()
        self.model_name = "LSTM"
        self.product_emb = nn.Embedding(num_products, embedding_dim, padding_idx=0)
        self.action_emb = nn.Embedding(num_actions, embedding_dim, padding_idx=0)
        input_dim = embedding_dim * 2
        self.lstm = nn.LSTM(input_dim, hidden_dim, num_layers=2, 
                           batch_first=True, dropout=0.2)
        self.dropout = nn.Dropout(0.2)
        self.fc_action = nn.Linear(hidden_dim, num_actions)
        self.fc_product = nn.Linear(hidden_dim, num_products)

    def forward(self, user_ids, product_ids, action_ids):
        p = self.product_emb(product_ids)
        a = self.action_emb(action_ids)
        x = torch.cat([p, a], dim=-1)
        out, (h, c) = self.lstm(x)
        last = self.dropout(out[:, -1, :])
        return self.fc_action(last), self.fc_product(last)'''
add_code_block(code_lstm)

# BiLSTM
add_heading_custom('3.1.3. BiLSTM (Bidirectional LSTM + Attention)', level=3)
doc.add_paragraph(
    'Mô hình BiLSTM sử dụng 2-layer Bidirectional LSTM với Attention mechanism. '
    'BiLSTM đọc sequence cả 2 chiều (forward + backward), kết hợp với Attention pooling '
    'để tập trung vào các timestep quan trọng nhất.'
)

code_bilstm = '''class BehaviorBiLSTM(nn.Module):
    def __init__(self, num_users, num_products, num_actions,
                 embedding_dim=64, hidden_dim=128):
        super().__init__()
        self.model_name = "BiLSTM"
        self.product_emb = nn.Embedding(num_products, embedding_dim, padding_idx=0)
        self.action_emb = nn.Embedding(num_actions, embedding_dim, padding_idx=0)
        input_dim = embedding_dim * 2
        self.bilstm = nn.LSTM(input_dim, hidden_dim, num_layers=2,
                              batch_first=True, dropout=0.2, bidirectional=True)
        self.dropout = nn.Dropout(0.2)
        self.attention = nn.Linear(hidden_dim * 2, 1)
        self.fc_action = nn.Linear(hidden_dim * 2, num_actions)
        self.fc_product = nn.Linear(hidden_dim * 2, num_products)

    def forward(self, user_ids, product_ids, action_ids):
        p = self.product_emb(product_ids)
        a = self.action_emb(action_ids)
        x = torch.cat([p, a], dim=-1)
        out, _ = self.bilstm(x)
        attn_weights = torch.softmax(self.attention(out), dim=1)
        context = (out * attn_weights).sum(dim=1)
        context = self.dropout(context)
        return self.fc_action(context), self.fc_product(context)'''
add_code_block(code_bilstm)

add_heading_custom('3.2. Hyperparameters', level=2)

table = doc.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Parameter'
hdr[1].text = 'Value'
params = [
    ('Embedding Dim', '64'),
    ('Hidden Dim', '128'),
    ('Max Epochs', '50'),
    ('Batch Size', '32'),
    ('Learning Rate', '0.003'),
    ('Sequence Length', '5'),
    ('Train/Test Split', '80% / 20%'),
    ('Optimizer', 'Adam (weight_decay=1e-4)'),
    ('LR Scheduler', 'CosineAnnealingLR'),
    ('Loss Function', 'Weighted CrossEntropy (action) + 0.3 * CrossEntropy (product)'),
    ('Early Stopping', 'patience=8 trên val F1-Score'),
    ('Gradient Clipping', 'max_norm=1.0'),
]
for name, val in params:
    row = table.add_row().cells
    row[0].text = name
    row[1].text = val

add_heading_custom('3.3. Training Code', level=2)
doc.add_paragraph('File: services/ai-new-service/app/training/trainer.py')

code_train = '''def _train_single(self, model, train_loader, test_loader, class_weights, epochs=50):
    optimizer = torch.optim.Adam(model.parameters(), lr=0.003, weight_decay=1e-4)
    scheduler = CosineAnnealingLR(optimizer, T_max=epochs, eta_min=1e-5)
    
    # Class-weighted loss cho imbalanced action distribution
    weight_tensor = torch.tensor(class_weights, dtype=torch.float32).to(self.device)
    criterion_action = nn.CrossEntropyLoss(weight=weight_tensor)
    criterion_product = nn.CrossEntropyLoss()
    
    for epoch in range(epochs):
        model.train()
        for user_ids, product_ids, action_ids, lbl_action, lbl_product in train_loader:
            optimizer.zero_grad()
            action_pred, product_pred = model(user_ids, product_ids, action_ids)
            loss = criterion_action(action_pred, lbl_action) + \\
                   0.3 * criterion_product(product_pred, lbl_product)
            loss.backward()
            nn.utils.clip_grad_norm_(model.parameters(), max_norm=1.0)
            optimizer.step()
        scheduler.step()
        
        # Early stopping on val F1
        val_f1 = f1_score(all_val_true, all_val_pred, average="weighted")
        if val_f1 > best_val_f1:
            best_state = model.state_dict()
        elif no_improve >= 8:
            break  # Early stopping'''
add_code_block(code_train)

add_heading_custom('3.4. Đánh giá và So sánh 3 mô hình', level=2)

add_heading_custom('3.4.1. Độ đo sử dụng', level=3)
metrics_desc = [
    'Accuracy: Tỷ lệ dự đoán đúng tổng thể',
    'Precision (weighted): Tỷ lệ dự đoán positive đúng cho mỗi class, có trọng số theo support',
    'Recall (weighted): Tỷ lệ tìm được positive thật cho mỗi class',
    'F1-Score (weighted): Trung bình điều hòa của Precision và Recall',
    'Confusion Matrix: Ma trận nhầm lẫn cho từng class',
]
for m in metrics_desc:
    doc.add_paragraph(m, style='List Bullet')

add_heading_custom('3.4.2. Bảng kết quả so sánh', level=3)
doc.add_paragraph('[GHI CHÚ: Thay số liệu thật sau khi chạy training]')

table = doc.add_table(rows=1, cols=7)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
headers = ['Model', 'Accuracy', 'Precision', 'Recall', 'F1-Score', 'Product Acc', 'Time']
for i, h in enumerate(headers):
    hdr[i].text = h
# Sample rows
sample_results = [
    ('RNN', 'XX.XX%', 'XX.XX%', 'XX.XX%', 'XX.XX%', 'XX.XX%', 'XXs'),
    ('LSTM', 'XX.XX%', 'XX.XX%', 'XX.XX%', 'XX.XX%', 'XX.XX%', 'XXs'),
    ('BiLSTM ★', 'XX.XX%', 'XX.XX%', 'XX.XX%', 'XX.XX%', 'XX.XX%', 'XXs'),
]
for vals in sample_results:
    row = table.add_row().cells
    for i, v in enumerate(vals):
        row[i].text = v

p = doc.add_paragraph()
run = p.add_run('[Thay bằng kết quả thật từ: http://localhost:8009/api/ai-new/train/status]')
run.font.color.rgb = RGBColor(255, 0, 0)
run.italic = True
run.font.size = Pt(10)

add_heading_custom('3.4.3. Visualization — Biểu đồ Training Curves', level=3)
plots_dir = r'C:\WORK\ThayQue\FINAL\Ecomerce-Dijango\services\ai-new-service\static\plots'
add_image_or_placeholder(
    os.path.join(plots_dir, 'training_curves.png'),
    'Hình 1: Training & Validation Loss/Accuracy qua các epoch cho 3 mô hình'
)

add_heading_custom('3.4.4. Visualization — So sánh Metrics', level=3)
add_image_or_placeholder(
    os.path.join(plots_dir, 'metrics_comparison.png'),
    'Hình 2: So sánh Accuracy, Precision, Recall, F1-Score giữa RNN, LSTM, BiLSTM'
)

add_heading_custom('3.4.5. Visualization — Confusion Matrices', level=3)
add_image_or_placeholder(
    os.path.join(plots_dir, 'confusion_matrices.png'),
    'Hình 3: Confusion Matrix cho 3 mô hình'
)

add_heading_custom('3.4.6. Visualization — Per-Class F1-Score', level=3)
add_image_or_placeholder(
    os.path.join(plots_dir, 'per_class_f1.png'),
    'Hình 4: F1-Score theo từng loại action cho 3 mô hình'
)

add_heading_custom('3.4.7. Visualization — Radar Chart', level=3)
add_image_or_placeholder(
    os.path.join(plots_dir, 'radar_chart.png'),
    'Hình 5: Radar Chart tổng hợp hiệu năng 3 mô hình'
)

add_heading_custom('3.5. Đánh giá và lựa chọn model_best', level=2)
doc.add_paragraph(
    'Dựa trên kết quả đánh giá, chúng tôi chọn mô hình có F1-Score cao nhất làm model_best. '
    'Lý do chọn F1-Score thay vì Accuracy:\n\n'
    '• F1-Score cân bằng giữa Precision và Recall, phù hợp với bài toán multi-class imbalanced.\n'
    '• Dữ liệu behavior không cân bằng (view chiếm ~30%, purchase chỉ ~5%), '
    'nên Accuracy đơn thuần có thể bị misleading.\n'
    '• Class-weighted loss giúp model học tốt cả các class thiểu số.\n\n'
    '[GHI CHÚ: Viết thêm nhận xét cụ thể dựa trên kết quả thực tế sau khi train]'
)

doc.add_page_break()

# ════════════════════════════════════════════════
# 5. CÂU 2b — KB GRAPH (Neo4j)
# ════════════════════════════════════════════════
add_heading_custom('4. CÂU 2b: Knowledge Base Graph (Neo4j)', level=1)

add_heading_custom('4.1. Mô tả', level=2)
doc.add_paragraph(
    'Knowledge Base Graph được xây dựng trên Neo4j, biểu diễn quan hệ giữa User, Product và Category '
    'dựa trên dữ liệu behavior đã sinh ở Câu 1.'
)

add_heading_custom('4.2. Cấu trúc Graph', level=2)

add_heading_custom('Nodes:', level=3)
nodes = [
    'User — id (UUID)',
    'Product — id, name, price, brand, rating, description, stock',
    'Category — id, name',
]
for n in nodes:
    doc.add_paragraph(n, style='List Bullet')

add_heading_custom('Relationships:', level=3)
rels = [
    '(User)-[:VIEW]->(Product) — User đã xem sản phẩm, count = số lần',
    '(User)-[:CLICK]->(Product) — User đã click sản phẩm',
    '(User)-[:ADD_TO_CART]->(Product) — User đã thêm vào giỏ hàng',
    '(User)-[:PURCHASE]->(Product) — User đã mua sản phẩm',
    '(User)-[:SEARCH]->(Product) — User đã tìm kiếm',
    '(User)-[:WISHLIST]->(Product) — User đã thêm vào wishlist',
    '(User)-[:REVIEW]->(Product) — User đã đánh giá',
    '(Product)-[:BELONGS_TO]->(Category) — Sản phẩm thuộc danh mục',
    '(Product)-[:CO_PURCHASED]->(Product) — Được mua cùng nhau (≥2 users)',
    '(Product)-[:CO_VIEWED]->(Product) — Được xem cùng nhau (≥3 users)',
]
for r in rels:
    doc.add_paragraph(r, style='List Bullet')

add_heading_custom('4.3. Code xây dựng Graph', level=2)
doc.add_paragraph('File: services/ai-new-service/app/knowledge_base/kb_graph.py')

code_neo4j = '''class KnowledgeBaseGraph:
    async def build_graph(self):
        # 1. Create constraints
        session.run("CREATE CONSTRAINT IF NOT EXISTS FOR (u:User) REQUIRE u.id IS UNIQUE")
        session.run("CREATE CONSTRAINT IF NOT EXISTS FOR (p:Product) REQUIRE p.id IS UNIQUE")
        
        # 2. Create Product + Category nodes
        for p in products:
            session.run("""MERGE (pr:Product {id: $id})
                          SET pr.name=$name, pr.price=$price, pr.brand=$brand""")
            session.run("""MATCH (pr:Product {id:$pid}), (c:Category {id:$cid})
                          MERGE (pr)-[:BELONGS_TO]->(c)""")
        
        # 3. Create User nodes + behavior edges
        for uid, product_actions in user_actions.items():
            session.run("MERGE (u:User {id: $id})", id=uid)
            for pid, actions in product_actions.items():
                for action in set(actions):
                    session.run(f"""MATCH (u:User {{id:$uid}}), (p:Product {{id:$pid}})
                                   MERGE (u)-[r:{action.upper()}]->(p)
                                   SET r.count = $count""")
        
        # 4. Co-purchase edges
        session.run("""MATCH (u:User)-[:PURCHASE]->(p1:Product),
                             (u)-[:PURCHASE]->(p2:Product)
                       WHERE p1.id < p2.id
                       WITH p1, p2, COUNT(u) AS co_count WHERE co_count >= 2
                       MERGE (p1)-[r:CO_PURCHASED]->(p2) SET r.count = co_count""")'''
add_code_block(code_neo4j)

add_heading_custom('4.4. Mẫu 20 dòng data trong Graph', level=2)
doc.add_paragraph('Cypher query: MATCH (u:User)-[r]->(p:Product) RETURN u.id, type(r), p.name, r.count LIMIT 20')

table = doc.add_table(rows=1, cols=4)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'User ID'
hdr[1].text = 'Relationship'
hdr[2].text = 'Product Name'
hdr[3].text = 'Count'

sample_rels = ['VIEW', 'CLICK', 'ADD_TO_CART', 'PURCHASE', 'VIEW', 'SEARCH',
               'VIEW', 'CLICK', 'WISHLIST', 'VIEW', 'PURCHASE', 'REVIEW',
               'VIEW', 'CLICK', 'ADD_TO_CART', 'VIEW', 'SEARCH', 'CLICK',
               'VIEW', 'PURCHASE']
for i in range(20):
    row = table.add_row().cells
    row[0].text = f'user-uuid-...{i%5+1}'
    row[1].text = sample_rels[i]
    row[2].text = f'[Tên sản phẩm {i+1}]'
    row[3].text = str((i % 5) + 1)

p = doc.add_paragraph()
run = p.add_run('[Thay bằng data thật từ Neo4j Browser: http://localhost:7474]')
run.font.color.rgb = RGBColor(255, 0, 0)
run.italic = True
run.font.size = Pt(10)

add_heading_custom('4.5. Ảnh Graph từ Neo4j', level=2)
add_image_or_placeholder('', 'Hình 6: Knowledge Base Graph trên Neo4j Browser (MATCH (n) RETURN n LIMIT 100)')

p = doc.add_paragraph()
run = p.add_run('[CHÈN ẢNH: Mở http://localhost:7474 → chạy query MATCH (n)-[r]-(m) RETURN n,r,m LIMIT 200 → chụp ảnh graph]')
run.font.color.rgb = RGBColor(255, 0, 0)
run.italic = True

doc.add_page_break()

# ════════════════════════════════════════════════
# 6. CÂU 2c — RAG + CHAT
# ════════════════════════════════════════════════
add_heading_custom('5. CÂU 2c: Xây dựng RAG và Chat dựa trên KB_Graph', level=1)

add_heading_custom('5.1. Kiến trúc RAG', level=2)
doc.add_paragraph(
    'Hệ thống RAG (Retrieval-Augmented Generation) kết hợp 2 nguồn context:\n\n'
    '1. Knowledge Graph (Neo4j): Lấy profile hành vi user, sản phẩm tương tự, '
    'collaborative filtering qua graph traversal\n'
    '2. Vector Store (ChromaDB): Index tất cả sản phẩm bằng sentence-transformers '
    '(model: paraphrase-multilingual-MiniLM-L12-v2), tìm kiếm semantic similarity\n\n'
    'Flow xử lý:\n'
    'User Message → Intent Detection → Query KB Graph + Search Vector Store → Generate Response'
)

add_heading_custom('5.2. Intent Detection', level=2)
intents = [
    'Gợi ý/Recommend: "gợi ý", "recommend", "nên mua", "tư vấn"',
    'So sánh: "so sánh", "compare", "khác nhau", "nào tốt hơn"',
    'Giá cả: "giá", "price", "bao nhiêu", "rẻ", "đắt"',
    'Tìm kiếm: "tìm", "search", "kiếm", "có bán"',
    'General: Các câu hỏi khác',
]
for intent in intents:
    doc.add_paragraph(intent, style='List Bullet')

add_heading_custom('5.3. Code RAG Chat Engine', level=2)
doc.add_paragraph('File: services/ai-new-service/app/rag/chat_engine.py')

code_rag = '''class RAGChatEngine:
    def chat(self, user_message, user_id=None, product_id=None):
        # 1. Get context from Knowledge Graph
        graph_context = self.kb_graph.graph_context_for_rag(
            user_id=user_id, product_id=product_id
        )
        
        # 2. Search products in Vector Store (ChromaDB)
        search_results = self.search_products(user_message, n_results=5)
        vector_context = ""
        for item in search_results:
            meta = item["metadata"]
            vector_context += f"- {meta['name']} ({meta['brand']}) - {meta['price']}đ\\n"
        
        # 3. Detect intent and generate response
        if "gợi ý" in query or "recommend" in query:
            return self._recommend_response(query, graph_context, vector_context)
        elif "so sánh" in query:
            return self._compare_response(query, vector_context)
        elif "giá" in query:
            return self._price_response(query, vector_context)
        else:
            return self._general_response(query, graph_context, vector_context)'''
add_code_block(code_rag)

add_heading_custom('5.4. API Endpoints', level=2)

table = doc.add_table(rows=1, cols=3)
table.style = 'Light Grid Accent 1'
hdr = table.rows[0].cells
hdr[0].text = 'Method'
hdr[1].text = 'Endpoint'
hdr[2].text = 'Mô tả'
endpoints = [
    ('POST', '/api/ai-new/chat', 'Chat với RAG (gửi message, nhận reply)'),
    ('GET', '/api/ai-new/recommend/{user_id}', 'Lấy recommendations cho user'),
    ('POST', '/api/ai-new/track', 'Ghi nhận hành vi user'),
    ('POST', '/api/ai-new/train', 'Trigger training model'),
    ('GET', '/api/ai-new/train/status', 'Xem kết quả training'),
    ('POST', '/api/ai-new/kb/build', 'Rebuild Knowledge Graph'),
    ('GET', '/api/ai-new/stats', 'Thống kê behavior data'),
]
for method, ep, desc in endpoints:
    row = table.add_row().cells
    row[0].text = method
    row[1].text = ep
    row[2].text = desc

add_heading_custom('5.5. Ảnh minh họa Chat', level=2)
add_image_or_placeholder('', 'Hình 7: Giao diện Chat AI (RAG) trên frontend')

doc.add_page_break()

# ════════════════════════════════════════════════
# 7. CÂU 2d — TRIỂN KHAI TÍCH HỢP
# ════════════════════════════════════════════════
add_heading_custom('6. CÂU 2d: Triển khai tích hợp trong Hệ E-commerce', level=1)

add_heading_custom('6.1. Kiến trúc tích hợp', level=2)
doc.add_paragraph(
    'AI New Service được tích hợp vào hệ thống e-commerce microservices thông qua:\n\n'
    '• Docker Compose: Thêm 3 containers mới (postgres-ai-new, neo4j, ai-new-service)\n'
    '• API Gateway: Thêm route "ai-new" vào SERVICE_MAP → proxy đến ai-new-service:8000\n'
    '• Frontend React: Thêm page AINewRecommend.jsx với 3 tabs\n'
    '• Behavior Tracking: Cập nhật aiTracking.js gửi event đến cả 2 AI services\n\n'
    'Lưu ý: KHÔNG thay đổi logic của các service cũ.'
)

add_heading_custom('6.2. Docker Compose', level=2)
code_docker = '''# docker-compose.yml (phần thêm mới)
postgres-ai-new:
  image: postgres:15-alpine
  ports: ["5439:5432"]
  environment:
    POSTGRES_DB: ai_new_db

neo4j:
  image: neo4j:5-community
  ports: ["7474:7474", "7687:7687"]
  environment:
    NEO4J_AUTH: neo4j/neo4jpassword

ai-new-service:
  build: ./services/ai-new-service
  ports: ["8009:8000"]
  depends_on: [postgres-ai-new, neo4j, product-service, redis]'''
add_code_block(code_docker)

add_heading_custom('6.3. Gateway Integration', level=2)
doc.add_paragraph('File: gateway/api_gateway/views.py')
code_gw = '''SERVICE_MAP = {
    ...
    'ai': 'http://ai-service:8000',
    'ai-new': 'http://ai-new-service:8000',  # NEW
}'''
add_code_block(code_gw)

add_heading_custom('6.4. Frontend — Tab AI Recommend', level=2)
doc.add_paragraph(
    'Thêm trang mới tại route /ai-recommend với 3 tabs:\n'
    '1. Gợi ý sản phẩm: Hiển thị danh sách sản phẩm recommend (ảnh, tên, giá, rating, brand)\n'
    '2. Chat AI (RAG): Giao diện chat với RAG engine\n'
    '3. Thống kê & Model: Dashboard hiển thị behavior stats và training results'
)
doc.add_paragraph('File: frontend/src/pages/AINewRecommend.jsx')

add_heading_custom('6.5. Behavior Tracking Integration', level=2)
doc.add_paragraph(
    'Khi user view sản phẩm hoặc add to cart, frontend gửi tracking event '
    'đến cả ai-service (cũ) và ai-new-service (mới) để cập nhật bảng UserBehaviorData.'
)
doc.add_paragraph('File: frontend/src/utils/aiTracking.js')

code_tracking = '''export function trackBehavior(userId, productId, eventType = 'view_detail') {
  // Track to original ai-service (giữ nguyên)
  axios.post(AI_TRACK_URL, { user_id, product_id, event_type }).catch(() => {});
  
  // Also track to ai-new-service (bảng UserBehaviorData)
  const actionMap = { view_detail: 'view', add_to_cart: 'add_to_cart', click: 'click' };
  axios.post(AI_NEW_TRACK_URL, {
    user_id: String(userId),
    product_id: String(productId),
    action: actionMap[eventType] || eventType,
  }).catch(() => {});
}'''
add_code_block(code_tracking)

add_heading_custom('6.6. Ảnh giao diện', level=2)

add_image_or_placeholder('', 'Hình 8: Tab "AI Recommend" trên Header navigation')
add_image_or_placeholder('', 'Hình 9: Giao diện Gợi ý sản phẩm (hiển thị ảnh, tên, giá, rating)')
add_image_or_placeholder('', 'Hình 10: Giao diện Chat AI với user')
add_image_or_placeholder('', 'Hình 11: Giao diện Thống kê & Model Training Status')

add_heading_custom('6.7. Hướng dẫn chạy', level=2)
doc.add_paragraph(
    'Chạy toàn bộ hệ thống:\n\n'
    '    cd Ecomerce-Dijango\n'
    '    docker compose up --build\n\n'
    'Sau khi khởi động (~5-10 phút lần đầu):\n'
    '• Frontend: http://localhost:3000 → click tab "AI Recommend"\n'
    '• AI New Service API: http://localhost:8009\n'
    '• Neo4j Browser: http://localhost:7474 (neo4j/neo4jpassword)\n'
    '• Xem plots: http://localhost:8009/api/ai-new/plots/training_curves.png\n'
    '• Xem report: http://localhost:8009/api/ai-new/report'
)

# ── Save
output_path = r'C:\WORK\ThayQue\FINAL\Ecomerce-Dijango\aiservice02_baitap.docx'
doc.save(output_path)
print(f'Saved to: {output_path}')
